"""Policy document parser, text extractor, and metadata detector."""
import os
import re
import uuid
from pathlib import Path
from typing import Tuple, Dict, Any, Optional
from datetime import datetime

import pypdf
import docx

ALLOWED_EXTENSIONS = {".txt", ".pdf", ".docx"}
MAX_FILE_SIZE_BYTES = 10 * 1024 * 1024  # 10 MB

SUPPORTED_CATEGORIES = [
    "Access Control",
    "Authentication / Password",
    "Account Management",
    "Data Protection",
    "Logging and Monitoring"
]

CATEGORY_KEYWORDS = {
    "Access Control": [
        "access control", "privileged", "role-based", "rbac", "least privilege",
        "access review", "authorization", "permission", "entitlement", "segregation of duties"
    ],
    "Authentication / Password": [
        "authentication", "password", "mfa", "multi-factor", "2fa", "credential",
        "authenticator", "sso", "single sign-on", "biometric", "passcode"
    ],
    "Account Management": [
        "account management", "provisioning", "deprovisioning", "leaver", "terminated",
        "dormant", "inactive account", "user lifecycle", "offboarding", "disable account"
    ],
    "Data Protection": [
        "data protection", "encryption", "dlp", "data loss prevention", "data classification",
        "confidentiality", "pii", "at rest", "in transit", "cryptographic"
    ],
    "Logging and Monitoring": [
        "logging", "monitoring", "audit logs", "audit trail", "siem", "telemetry",
        "log retention", "incident detection", "soc", "syslog"
    ]
}


def validate_file(filename: str, file_size: int) -> Tuple[bool, str]:
    """Validate file extension and size safely."""
    ext = Path(filename).suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        return False, f"Unsupported file format '{ext}'. Allowed formats: PDF, DOCX, TXT."
    if file_size > MAX_FILE_SIZE_BYTES:
        return False, f"File size exceeds 10 MB limit ({file_size / (1024*1024):.2f} MB)."
    return True, ""


def extract_text_from_file(file_path: Path) -> str:
    """Extract raw text from TXT, PDF, or DOCX safely without executing."""
    ext = file_path.suffix.lower()
    text = ""
    if ext == ".txt":
        try:
            text = file_path.read_text(encoding="utf-8")
        except UnicodeDecodeError:
            text = file_path.read_text(encoding="latin-1", errors="ignore")
    elif ext == ".pdf":
        reader = pypdf.PdfReader(str(file_path))
        extracted = []
        for page in reader.pages:
            t = page.extract_text()
            if t:
                extracted.append(t)
        text = "\n".join(extracted)
    elif ext == ".docx":
        doc = docx.Document(str(file_path))
        extracted = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    extracted.append(row_text)
        text = "\n".join(extracted)
    else:
        raise ValueError(f"Unsupported format: {ext}")
    return text.strip()


def detect_metadata(text: str, filename: str) -> Dict[str, Any]:
    """Detect policy metadata deterministically from document content."""
    lines = [line.strip() for line in text.splitlines() if line.strip()]

    # 1. Policy Name
    policy_name = "Not detected"
    name_patterns = [
        r"(?:Policy\s*Name|Title|Document\s*Title)\s*[:=|]?\s*([^\n\r|]+)",
        r"^#\s*([^\n\r]+)",
        r"^([A-Z][A-Za-z0-9\s\-–—]{4,80}\s*Policy)$"
    ]
    for pattern in name_patterns:
        match = re.search(pattern, text, re.IGNORECASE | re.MULTILINE)
        if match:
            raw_name = match.group(1).strip()
            # Clean trailing metadata labels if on the same line
            clean_name = re.split(r"(?:Category|Version|Owner|Effective|Review|Date)\s*[:=|]", raw_name, flags=re.IGNORECASE)[0].strip()
            if clean_name:
                policy_name = clean_name
                break
    if policy_name == "Not detected" and lines:
        first_line = lines[0]
        if len(first_line) < 80 and not first_line.lower().startswith("version") and not first_line.lower().startswith("owner"):
            policy_name = first_line
        else:
            # Fallback to sanitized filename
            clean_name = Path(filename).stem.replace("_", " ").replace("-", " ").title()
            policy_name = clean_name

    # 2. Version
    version = "Not detected"
    ver_match = re.search(r"(?:Version|Ver\.?|v)\s*[:=|]?\s*([0-9]+(?:\.[0-9]+)*)", text, re.IGNORECASE)
    if ver_match:
        version = f"v{ver_match.group(1)}"

    # 3. Owner
    owner = "Not detected"
    owner_match = re.search(r"(?:^|[\n\r|])\s*(?:Owner|Document\s*Owner|Author|Maintained\s*by|Approved\s*by)\b\s*[:=|]?\s*([^\n\r|]+)", text, re.IGNORECASE)
    if owner_match:
        raw_owner = owner_match.group(1).strip()
        # Clean trailing labels
        raw_owner = re.split(r"(?:Effective|Review|Date|Version|Category)", raw_owner, flags=re.IGNORECASE)[0].strip()
        if raw_owner:
            owner = raw_owner

    # 4. Effective Date
    effective_date = "Not detected"
    eff_match = re.search(r"(?:Effective\s*Date|Issued|Date\s*of\s*Issue|Effective)\s*[:=|]?\s*([0-9]{4}[-/][0-9]{1,2}[-/][0-9]{1,2}|[A-Za-z]+\s+[0-9]{1,2},?\s+[0-9]{4})", text, re.IGNORECASE)
    if eff_match:
        effective_date = eff_match.group(1).strip()

    # 5. Review Date
    review_date = "Not detected"
    rev_match = re.search(r"(?:Review\s*Date|Next\s*Review|Renewal\s*Date|Valid\s*Until)\s*[:=|]?\s*([0-9]{4}[-/][0-9]{1,2}[-/][0-9]{1,2}|[A-Za-z]+\s+[0-9]{1,2},?\s+[0-9]{4})", text, re.IGNORECASE)
    if rev_match:
        review_date = rev_match.group(1).strip()

    # 6. Policy Category
    category = "Policy category could not be confidently determined."
    # Check explicit Category field first
    cat_match = re.search(r"Category\s*[:=|]?\s*([^\n\r|]+)", text, re.IGNORECASE)
    if cat_match:
        raw_cat = cat_match.group(1).strip()
        raw_cat = re.split(r"(?:Version|Owner|Effective|Review|Date)\s*[:=|]", raw_cat, flags=re.IGNORECASE)[0].strip()
        for sc in SUPPORTED_CATEGORIES:
            if sc.lower() in raw_cat.lower() or raw_cat.lower() in sc.lower():
                category = sc
                break

    if category == "Policy category could not be confidently determined.":
        category_scores = {}
        lower_text = text.lower()
        for cat, kws in CATEGORY_KEYWORDS.items():
            score = sum(lower_text.count(kw) for kw in kws)
            category_scores[cat] = score

        sorted_scores = sorted(category_scores.items(), key=lambda x: x[1], reverse=True)
        best_cat, best_score = sorted_scores[0]
        second_cat, second_score = sorted_scores[1]

        if best_score > 0 and not (best_score == second_score and best_score < 2):
            category = best_cat

    return {
        "name": policy_name,
        "version": version,
        "owner": owner,
        "effective_date": effective_date,
        "review_date": review_date,
        "category": category,
        "word_count": len(text.split())
    }
