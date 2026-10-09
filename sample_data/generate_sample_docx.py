"""Generate sample policy DOCX files for testing PolicyGuard document upload."""
from pathlib import Path
import docx

SAMPLE_DIR = Path(__file__).resolve().parent.parent / "sample_data"
SAMPLE_DIR.mkdir(parents=True, exist_ok=True)

def build_docx(filename: str, metadata: dict, sections: list):
    doc = docx.Document()
    
    # Header Title
    doc.add_heading(metadata["Policy Name"], level=0)
    
    # Metadata Table
    table = doc.add_table(rows=6, cols=2)
    table.style = 'Light Shading Accent 1'
    fields = [
        ("Policy Name", metadata["Policy Name"]),
        ("Category", metadata["Category"]),
        ("Version", metadata["Version"]),
        ("Owner", metadata["Owner"]),
        ("Effective Date", metadata["Effective Date"]),
        ("Review Date", metadata["Review Date"])
    ]
    for i, (k, v) in enumerate(fields):
        row = table.rows[i]
        row.cells[0].text = k
        row.cells[1].text = v
    
    doc.add_paragraph()  # Spacing
    
    # Content Sections
    for heading, paragraphs in sections:
        doc.add_heading(heading, level=1)
        for p in paragraphs:
            doc.add_paragraph(p)
            
    out_path = SAMPLE_DIR / filename
    doc.save(str(out_path))
    print(f"Generated DOCX: {out_path}")
    return out_path

def main():
    # 1. Access Control Policy DOCX
    build_docx(
        "sample_access_control_policy.docx",
        {
            "Policy Name": "Enterprise Access Control & Identity Security Policy",
            "Version": "3.2",
            "Owner": "Identity and Access Management Governance Team",
            "Effective Date": "2026-01-01",
            "Review Date": "2026-12-31",
            "Category": "Access Control"
        },
        [
            ("1. Purpose & Strategic Scope", [
                "This policy establishes mandatory identity governance, authentication safeguards, and account lifecycle controls across all cloud platforms, infrastructure consoles, and corporate information systems.",
                "All internal employees, contractors, and administrative operators must strictly adhere to these controls to safeguard organizational assets."
            ]),
            ("2. Measurable Policy Directives", [
                "- Privileged accounts must use multi-factor authentication (MFA) prior to gaining access to critical environments.",
                "- Administrative access must be reviewed every 90 days by designated security governance custodians.",
                "- Terminated users must have access disabled within 24 hours of official HR separation notification.",
                "- Enabled accounts without activity for more than 45 days must be reviewed and remediated to prevent dormant credential exploitation."
            ]),
            ("3. Continuous Compliance & Enforcement", [
                "Adherence to these technical directives will be assessed continuously by automated telemetry checks. Violations are logged as risk findings and assigned to resource owners for mandatory remediation."
            ])
        ]
    )

if __name__ == "__main__":
    main()
